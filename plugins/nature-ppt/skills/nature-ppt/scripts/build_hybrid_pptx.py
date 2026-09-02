#!/usr/bin/env python3
"""Build a responsive PowerPoint slide with a sharp background and live text."""

from __future__ import annotations

import argparse
import json
import math
import os
import tempfile
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt


EMU_PER_PT = 12700.0


def color_rgb(value: object) -> RGBColor:
    text = str(value or "#000000").lstrip("#")
    if len(text) == 3:
        text = "".join(char * 2 for char in text)
    if len(text) < 6:
        raise ValueError(f"unsupported color: {value}")
    return RGBColor(int(text[0:2], 16), int(text[2:4], 16), int(text[4:6], 16))


def finite(value: object, fallback: float = 0.0) -> float:
    try:
        number = float(value)
    except (TypeError, ValueError):
        return fallback
    return number if math.isfinite(number) else fallback


def load_manifest(path: Path | None) -> list[dict]:
    if not path:
        return []
    payload = json.loads(path.expanduser().resolve(strict=True).read_text(encoding="utf-8-sig"))
    if payload.get("schema_version") != "1.0" or not isinstance(payload.get("text_elements"), list):
        raise ValueError("text manifest must use schema_version 1.0 and contain text_elements")
    return payload["text_elements"]


def move_to_back(slide, shape) -> None:
    tree = slide.shapes._spTree
    tree.remove(shape._element)
    tree.insert(2, shape._element)


def add_live_text(slide, item: dict, image_box: tuple[float, float, float, float], source_size: tuple[int, int]) -> None:
    left, top, placed_width, placed_height = image_box
    source_width, source_height = source_size
    normalized = item.get("coordinate_space", "normalized") == "normalized"
    x = finite(item.get("x"))
    y = finite(item.get("y"))
    if normalized:
        x = left + x * placed_width
        y = top + y * placed_height
    else:
        x = left + x / source_width * placed_width
        y = top + y / source_height * placed_height
    raw_size = finite(item.get("font_size"), 12.0)
    if normalized and item.get("font_size_space") == "normalized":
        font_size = raw_size * placed_height
    else:
        font_size = raw_size / source_height * placed_height if not normalized else raw_size
    font_size = max(4.0, font_size)

    bbox = item.get("bbox_normalized") if isinstance(item.get("bbox_normalized"), dict) else None
    if bbox:
        box_left = left + finite(bbox.get("x")) * placed_width
        box_top = top + finite(bbox.get("y")) * placed_height
        box_width = max(font_size * 1.5, finite(bbox.get("width")) * placed_width)
        box_height = max(font_size * 1.35, finite(bbox.get("height")) * placed_height)
    else:
        contents = str(item.get("content", ""))
        box_width = max(font_size * 2.0, len(contents) * font_size * 0.72)
        box_height = font_size * 1.5
        anchor = str(item.get("text_anchor", "start")).lower()
        box_left = x - box_width / 2 if anchor == "middle" else x - box_width if anchor == "end" else x
        box_top = y - font_size * 1.08

    shape = slide.shapes.add_textbox(
        Emu(round(box_left * EMU_PER_PT)),
        Emu(round(box_top * EMU_PER_PT)),
        Emu(round(box_width * EMU_PER_PT)),
        Emu(round(box_height * EMU_PER_PT)),
    )
    shape.name = str(item.get("id") or f"NATURE_PPT_TEXT_{len(slide.shapes):04d}")
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = False
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    paragraph = frame.paragraphs[0]
    anchor = str(item.get("text_anchor", "start")).lower()
    paragraph.alignment = PP_ALIGN.CENTER if anchor == "middle" else PP_ALIGN.RIGHT if anchor == "end" else PP_ALIGN.LEFT
    run = paragraph.add_run()
    run.text = str(item.get("content", ""))
    font = run.font
    font.name = str(item.get("font_family") or "Arial")
    font.size = Pt(font_size)
    weight = str(item.get("font_weight", "normal")).lower()
    font.bold = weight == "bold" or (weight.isdigit() and int(weight) >= 600)
    font.italic = str(item.get("font_style", "normal")).lower() == "italic"
    font.color.rgb = color_rgb(item.get("fill"))
    shape.rotation = finite(item.get("rotation"), 0.0)


def build(
    background: Path,
    output: Path,
    manifest: Path | None = None,
    input_pptx: Path | None = None,
    slide_index: int = 0,
) -> dict:
    background = background.expanduser().resolve(strict=True)
    output = output.expanduser().resolve()
    with Image.open(background) as image:
        source_width, source_height = image.size
    prs = Presentation(str(input_pptx.expanduser().resolve(strict=True))) if input_pptx else Presentation()
    if not input_pptx:
        width_inches = 7.5 if source_height > source_width else 13.333
        height_inches = width_inches * source_height / source_width
        if height_inches > 13.333:
            height_inches = 13.333
            width_inches = height_inches * source_width / source_height
        prs.slide_width = Inches(width_inches)
        prs.slide_height = Inches(height_inches)
    if slide_index:
        if not 1 <= slide_index <= len(prs.slides):
            raise ValueError("slide-index is out of range")
        slide = prs.slides[slide_index - 1]
    else:
        slide = prs.slides.add_slide(prs.slide_layouts[6])

    slide_width = prs.slide_width / EMU_PER_PT
    slide_height = prs.slide_height / EMU_PER_PT
    scale = min(slide_width / source_width, slide_height / source_height)
    placed_width = source_width * scale
    placed_height = source_height * scale
    left = (slide_width - placed_width) / 2
    top = (slide_height - placed_height) / 2
    picture = slide.shapes.add_picture(
        str(background),
        Emu(round(left * EMU_PER_PT)),
        Emu(round(top * EMU_PER_PT)),
        Emu(round(placed_width * EMU_PER_PT)),
        Emu(round(placed_height * EMU_PER_PT)),
    )
    picture.name = "NATURE_PPT_BACKGROUND"
    move_to_back(slide, picture)
    items = load_manifest(manifest)
    for item in sorted(items, key=lambda value: int(value.get("paint_order", 0))):
        add_live_text(slide, item, (left, top, placed_width, placed_height), (source_width, source_height))

    output.parent.mkdir(parents=True, exist_ok=True)
    handle, temporary_name = tempfile.mkstemp(prefix=f".{output.stem}-", suffix=".pptx", dir=output.parent)
    os.close(handle)
    temporary = Path(temporary_name)
    try:
        prs.save(str(temporary))
        Presentation(str(temporary))
        os.replace(temporary, output)
    finally:
        temporary.unlink(missing_ok=True)
    return {
        "ok": True,
        "mode": "hybrid",
        "output_pptx": str(output),
        "slide_index": slide_index or len(prs.slides),
        "editable_text_count": len(items),
        "background_is_raster": True,
        "source_dimensions": [source_width, source_height],
    }


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--background-image", required=True, type=Path)
    parser.add_argument("--text-manifest", type=Path)
    parser.add_argument("--output-pptx", required=True, type=Path)
    parser.add_argument("--input-pptx", type=Path)
    parser.add_argument("--slide-index", type=int, default=0)
    args = parser.parse_args()
    report = build(args.background_image, args.output_pptx, args.text_manifest, args.input_pptx, args.slide_index)
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
