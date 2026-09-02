#!/usr/bin/env python3
"""Verify native and hybrid PPTX output plus isolated installation."""

from __future__ import annotations

import json
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.util import Inches


ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = ROOT / "plugins" / "nature-ppt" / "skills" / "nature-ppt" / "scripts"


def run(*args: object) -> subprocess.CompletedProcess[str]:
    return subprocess.run([str(value) for value in args], check=True, text=True, capture_output=True)


def main() -> int:
    with tempfile.TemporaryDirectory(prefix="nature-ppt-cross-") as raw:
        temp = Path(raw)
        cache = temp / "cache"
        run(sys.executable, SCRIPTS / "prepare_geometry_cache.py", "--input", ROOT / "tests" / "fixtures" / "editable.svg", "--output-dir", cache, "--job-id", "natureppt1")
        run(sys.executable, SCRIPTS / "cull_hidden_geometry.py", "--cache", cache / "geometry-cache.json", "--state", cache / "drawing-state.json")
        payload = json.loads((cache / "geometry-cache.json").read_text(encoding="utf-8"))
        assert payload["schema_version"] == 3 and any(atom["kind"] == "text" for atom in payload["atoms"])

        source = temp / "source.pptx"
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        sentinel = slide.shapes.add_textbox(Inches(0.1), Inches(0.1), Inches(1), Inches(0.3))
        sentinel.name = "PREEXISTING_SENTINEL"
        sentinel.text = "keep"
        prs.save(source)
        native = temp / "native.pptx"
        result = json.loads(run(sys.executable, SCRIPTS / "render_pptx_ooxml.py", "--geometry-cache", cache / "geometry-cache.json", "--input-pptx", source, "--output-pptx", native, "--slide-index", 1).stdout)
        assert result["native_object_count"] >= 4
        reopened = Presentation(native)
        names = [shape.name for shape in reopened.slides[0].shapes]
        texts = [shape.text for shape in reopened.slides[0].shapes if getattr(shape, "has_text_frame", False)]
        assert "PREEXISTING_SENTINEL" in names and "Nature PPT" in texts
        with zipfile.ZipFile(native) as package:
            assert not any(name.startswith("ppt/media/") for name in package.namelist())

        background = temp / "background.png"
        Image.new("RGB", (320, 180), "#14324a").save(background)
        hybrid = temp / "hybrid.pptx"
        hybrid_report = json.loads(run(sys.executable, SCRIPTS / "build_hybrid_pptx.py", "--background-image", background, "--text-manifest", ROOT / "tests" / "fixtures" / "nature-text-manifest.json", "--output-pptx", hybrid).stdout)
        assert hybrid_report["background_is_raster"] and hybrid_report["editable_text_count"] == 1
        hybrid_opened = Presentation(hybrid)
        assert any(shape.name == "NATURE_PPT_BACKGROUND" for shape in hybrid_opened.slides[-1].shapes)
        with zipfile.ZipFile(hybrid) as package:
            assert any(name.startswith("ppt/media/") for name in package.namelist())

        install_root = temp / "skills"
        existing = install_root / "nature-ppt"
        existing.mkdir(parents=True)
        (existing / "remote-backend.json").write_text('{"url":"https://example.invalid"}', encoding="utf-8")
        run(sys.executable, ROOT / "install.py", "--destination", install_root, "--force")
        assert (existing / "SKILL.md").is_file()
        assert json.loads((existing / "remote-backend.json").read_text(encoding="utf-8"))["url"].startswith("https://")
        assert [path.name for path in install_root.iterdir()] == ["nature-ppt"]

    print("CROSS_PLATFORM_OK|native=true|hybrid=true|existing_preserved=true|isolated_install=true")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
